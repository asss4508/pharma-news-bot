import base64
import io
import json
import math
import os
import re
import tempfile
from collections import defaultdict
from pathlib import Path

import requests as http_requests
import anthropic
from telegram import Update
from telegram.ext import Application, MessageHandler, CommandHandler, filters, ContextTypes

TELEGRAM_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN", "")
GITHUB_REPO = os.environ.get("GITHUB_REPO", "asss4508/pharma-news-bot")
TELEGRAM_API_ID = int(os.environ.get("TELEGRAM_API_ID", "0"))
TELEGRAM_API_HASH = os.environ.get("TELEGRAM_API_HASH", "")
TELETHON_SESSION = os.environ.get("TELETHON_SESSION", "")
VOYAGE_API_KEY = os.environ.get("VOYAGE_API_KEY", "")

INDEX_PATH = Path(__file__).parent.parent.parent / "data" / "index.json"
CHANNELS_FILE = "data/channels_to_sync.txt"

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".docx", ".txt", ".csv", ".md"}

# ── 텍스트 추출 ────────────────────────────────────────────────────────────────

def extract_from_bytes(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[시트: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(v) for v in row if v is not None)
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
        elif ext == ".docx":
            from docx import Document as DocxDoc
            doc = DocxDoc(io.BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            for enc in ("utf-8", "cp949", "euc-kr"):
                try:
                    return file_bytes.decode(enc)
                except UnicodeDecodeError:
                    continue
    except Exception as e:
        print(f"추출 오류 ({filename}): {e}")
    return ""

def split_chunks(text: str, source: str, chunk_size=800, overlap=150):
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    if not text:
        return []
    paragraphs = text.split('\n\n')
    chunks, buffer = [], ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(buffer) + len(para) > chunk_size:
            if buffer:
                chunks.append({"text": buffer.strip(), "source": source})
                buffer = buffer[-overlap:] + " " + para
            else:
                for i in range(0, len(para), chunk_size - overlap):
                    chunk = para[i:i + chunk_size]
                    if chunk.strip():
                        chunks.append({"text": chunk.strip(), "source": source})
                buffer = ""
        else:
            buffer = (buffer + "\n\n" + para).strip() if buffer else para
    if buffer.strip():
        chunks.append({"text": buffer.strip(), "source": source})
    return chunks

# ── BM25 ──────────────────────────────────────────────────────────────────────

def tokenize(text: str):
    return re.findall(r'[가-힣a-zA-Z0-9]+', text.lower())

class BM25:
    def __init__(self, chunks, k1=1.5, b=0.75):
        self.chunks = list(chunks)
        self.k1 = k1
        self.b = b
        self.n = len(self.chunks)
        self.avg_dl = 0
        self.df = defaultdict(int)
        self.tf = []
        total = 0
        for chunk in self.chunks:
            tokens = tokenize(chunk["text"])
            total += len(tokens)
            freq = defaultdict(int)
            for t in tokens:
                freq[t] += 1
            self.tf.append(freq)
            for t in freq:
                self.df[t] += 1
        self.avg_dl = total / self.n if self.n else 1

    def add_chunks(self, new_chunks):
        return BM25(self.chunks + new_chunks)

    def search(self, query: str, top_k=10):
        q_tokens = tokenize(query)
        scores = []
        for i, freq in enumerate(self.tf):
            dl = sum(freq.values())
            score = 0
            for t in q_tokens:
                if t not in freq:
                    continue
                df = self.df[t]
                idf = math.log((self.n - df + 0.5) / (df + 0.5) + 1)
                tf = freq[t]
                norm_tf = tf * (self.k1 + 1) / (tf + self.k1 * (1 - self.b + self.b * dl / self.avg_dl))
                score += idf * norm_tf
            if score > 0:
                scores.append((score, i))
        scores.sort(reverse=True)
        return [self.chunks[i] for _, i in scores[:top_k]]

# ── GitHub 저장 ────────────────────────────────────────────────────────────────

def _github_put(file_path: str, content_bytes: bytes, message: str) -> bool:
    if not GITHUB_TOKEN:
        return False
    headers = {
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github+json"
    }
    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{file_path}"
    existing = http_requests.get(url, headers=headers)
    try:
        sha = existing.json().get("sha") if existing.status_code == 200 else None
    except Exception:
        sha = None
    payload = {"message": message, "content": base64.b64encode(content_bytes).decode()}
    if sha:
        payload["sha"] = sha
    resp = http_requests.put(url, json=payload, headers=headers)
    return resp.status_code in (200, 201)

def save_to_github(filename: str, file_bytes: bytes) -> bool:
    # 원본 파일만 저장 → build_index.yml이 자동으로 인덱스 재빌드
    return _github_put(
        f"data/uploads/{filename}",
        file_bytes,
        f"feat: {filename} 업로드 (텔레그램)"
    )

def save_channel_text_to_github(channel: str, text: str) -> bool:
    # 채널 텍스트 파일 저장 → build_index.yml이 자동으로 인덱스 재빌드
    safe_name = re.sub(r'[^a-zA-Z0-9_]', '_', channel.lstrip('@'))
    return _github_put(
        f"data/channels/{safe_name}.txt",
        text.encode("utf-8"),
        f"chore: 채널 업데이트 ({channel})"
    )

def add_channel_to_github(channel: str):
    if not GITHUB_TOKEN:
        return False
    headers = {
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github+json"
    }
    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{CHANNELS_FILE}"
    existing = http_requests.get(url, headers=headers)
    if existing.status_code == 200:
        try:
            data = existing.json()
            sha = data.get("sha")
            content = base64.b64decode(data.get("content", "")).decode("utf-8")
        except Exception:
            sha = None
            content = "# 동기화할 텔레그램 채널 목록\n"
    else:
        sha = None
        content = "# 동기화할 텔레그램 채널 목록\n"

    normalized = "@" + channel.lstrip("@")
    if normalized in content or channel in content:
        return "already_exists"

    content = content.rstrip("\n") + f"\n{normalized}\n"
    payload = {
        "message": f"feat: 채널 추가 ({normalized})",
        "content": base64.b64encode(content.encode("utf-8")).decode(),
    }
    if sha:
        payload["sha"] = sha
    resp = http_requests.put(url, json=payload, headers=headers)
    return resp.status_code in (200, 201)

# ── Jina Reader (Notion/URL 텍스트 추출) ──────────────────────────────────────

def fetch_url_content(url: str) -> str:
    jina_url = f"https://r.jina.ai/{url}"
    try:
        resp = http_requests.get(jina_url, timeout=30, headers={"Accept": "text/plain"})
        if resp.status_code == 200:
            return resp.text
    except Exception as e:
        print(f"Jina 추출 오류: {e}")
    return ""

# ── Telethon 채널 동기화 ───────────────────────────────────────────────────────

async def sync_single_channel(channel: str) -> tuple[int, str]:
    if not (TELEGRAM_API_ID and TELEGRAM_API_HASH and TELETHON_SESSION):
        return 0, "Telethon 설정이 없습니다."
    try:
        from telethon import TelegramClient
        from telethon.sessions import StringSession
        client = TelegramClient(StringSession(TELETHON_SESSION), TELEGRAM_API_ID, TELEGRAM_API_HASH)
        await client.connect()
        if not await client.is_user_authorized():
            await client.disconnect()
            return 0, "Telethon 인증 실패"
        entity = await client.get_entity(channel)
        messages = []
        from datetime import datetime, timezone
        today = datetime.now(timezone.utc).date()
        async for msg in client.iter_messages(entity, limit=None):
            if not msg.date:
                continue
            msg_date = msg.date.astimezone(timezone.utc).date()
            if msg_date == today:
                if msg.text and msg.text.strip():
                    messages.append(msg.text.strip())
            elif msg_date < today:
                break
        await client.disconnect()
        if not messages:
            return 0, "오늘 메시지 없음"
        text = "\n\n".join(reversed(messages))
        return len(messages), text
    except Exception as e:
        return 0, str(e)

# ── Voyage AI 리랭킹 ──────────────────────────────────────────────────────────

def rerank_chunks(query: str, candidates: list, top_k: int = 5) -> list:
    if not VOYAGE_API_KEY or not candidates:
        return candidates[:top_k]
    try:
        import voyageai
        vo = voyageai.Client(api_key=VOYAGE_API_KEY)
        docs = [c["text"] for c in candidates]
        result = vo.rerank(query, docs, model="rerank-2-lite", top_k=min(top_k, len(docs)))
        return [candidates[r.index] for r in result.results]
    except Exception as e:
        print(f"Voyage rerank 오류: {e}")
        return candidates[:top_k]

# ── Claude 답변 ────────────────────────────────────────────────────────────────

def answer(question: str, bm25: BM25) -> str:
    # BM25로 후보 20개 → Voyage AI로 의미 기반 재정렬 → 상위 5개로 Claude 답변
    candidates = bm25.search(question, top_k=20)
    results = rerank_chunks(question, candidates, top_k=5)
    if not results:
        context = "관련 자료를 찾을 수 없습니다."
    else:
        context = "\n\n---\n\n".join(
            f"[출처: {r['source']}]\n{r['text']}" for r in results
        )
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    prompt = f"""다음은 참고 자료입니다:

{context}

---

위 자료를 바탕으로 질문에 답해줘.
- 자료에 근거한 내용만 답해줘. 자료에 없으면 명확히 없다고 말해줘.
- 출처(파일명 또는 채널명)를 함께 표시해줘.
- 한국어로 답해줘.
- 첫 문단에서 질문에 대한 결론을 먼저 말하고, 이어서 필요한 근거를 짧은 문단으로 설명해줘.
- 차분한 존댓말을 쓰고 과장, 반복, 불필요한 감탄사나 장식용 이모지는 피해주세요.
- 수치와 기준일을 함께 제시하고, 과거 자료의 내용을 현재 사실로 단정하지 마.
- 출처는 관련 설명 가까이에 표시하고, 자료에 없는 전망이나 투자 판단은 만들지 마.
- **나 __ 같은 마크다운 강조 기호는 절대 사용하지 마. 텍스트만 써.
- ## ### 같은 제목 기호도 쓰지 마.
- 일반 텍스트로만 작성해줘.

질문: {question}"""
    msg = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=2000,
        messages=[{"role": "user", "content": prompt}]
    )
    text = msg.content[0].text.strip()
    text = re.sub(r'\*{2,}', '', text)
    text = re.sub(r'#{1,6}\s*(.+)', r'◆ \1', text)
    text = re.sub(r'_{1,2}([^_\n]+)_{1,2}', r'\1', text)
    text = re.sub(r'^>\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

# ── Telegram 핸들러 ────────────────────────────────────────────────────────────

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if not doc:
        return

    filename = doc.file_name or "unknown"
    ext = Path(filename).suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        await update.message.reply_text(
            f"❌ 지원하지 않는 형식입니다.\n지원: PDF, Excel, Word, TXT, CSV"
        )
        return

    if doc.file_size and doc.file_size > 20 * 1024 * 1024:
        await update.message.reply_text("❌ 파일 크기가 20MB를 초과합니다.")
        return

    status = await update.message.reply_text(f"📥 {filename} 처리 중...")

    try:
        tg_file = await context.bot.get_file(doc.file_id)
        file_bytes = bytes(await tg_file.download_as_bytearray())

        text = extract_from_bytes(file_bytes, filename)
        if not text.strip():
            await status.edit_text(f"❌ {filename}에서 텍스트를 추출할 수 없습니다.")
            return

        new_chunks = split_chunks(text, filename)
        bm25: BM25 = context.bot_data["bm25"]
        context.bot_data["bm25"] = bm25.add_chunks(new_chunks)

        saved = save_to_github(filename, file_bytes)
        persist_msg = "💾 GitHub 저장 완료 (인덱스 자동 재빌드 중)" if saved else "⚠️ 이번 세션에만 유효"

        await status.edit_text(
            f"✅ {filename} 처리 완료!\n"
            f"📊 {len(new_chunks)}개 청크 추가\n"
            f"📚 총 {context.bot_data['bm25'].n}개 청크\n"
            f"{persist_msg}\n\n"
            f"이제 이 파일 내용을 질문할 수 있습니다."
        )

    except Exception as e:
        await status.edit_text(f"❌ {type(e).__name__}: {str(e)[:150]}")

async def handle_notion_url(update: Update, context: ContextTypes.DEFAULT_TYPE, url: str):
    status = await update.message.reply_text(f"📥 노션 페이지 가져오는 중...")
    try:
        text = fetch_url_content(url)
        if not text.strip():
            await status.edit_text("❌ 노션 페이지에서 텍스트를 가져올 수 없습니다.")
            return

        source = re.sub(r'https?://', '', url)[:60]
        new_chunks = split_chunks(text, source)
        bm25: BM25 = context.bot_data["bm25"]
        context.bot_data["bm25"] = bm25.add_chunks(new_chunks)

        if GITHUB_TOKEN:
            safe_name = re.sub(r'[^a-zA-Z0-9_]', '_', source)[:60]
            _github_put(
                f"data/uploads/{safe_name}.txt",
                text.encode("utf-8"),
                f"feat: URL 저장 ({source[:50]})"
            )
            persist_msg = "💾 GitHub 저장 완료 (인덱스 자동 재빌드 중)"
        else:
            persist_msg = "⚠️ 이번 세션에만 유효"

        await status.edit_text(
            f"✅ 노션 페이지 저장 완료!\n"
            f"📊 {len(new_chunks)}개 청크 추가\n"
            f"📚 총 {context.bot_data['bm25'].n}개 청크\n"
            f"{persist_msg}\n\n"
            f"이제 이 페이지 내용을 질문할 수 있습니다."
        )
    except Exception as e:
        await status.edit_text(f"❌ {type(e).__name__}: {str(e)[:150]}")

async def handle_telegram_channel(update: Update, context: ContextTypes.DEFAULT_TYPE, channel: str):
    status = await update.message.reply_text(f"📥 @{channel} 채널 동기화 중...")
    try:
        msg_count, text = await sync_single_channel(channel)
        if msg_count == 0:
            await status.edit_text(f"❌ 채널 동기화 실패: {text}")
            return

        source = f"@{channel}"
        new_chunks = split_chunks(text, source)
        bm25: BM25 = context.bot_data["bm25"]
        context.bot_data["bm25"] = bm25.add_chunks(new_chunks)

        if GITHUB_TOKEN:
            # 채널 텍스트 파일 저장 → build_index.yml이 인덱스 자동 재빌드
            save_channel_text_to_github(channel, text)
            # channels_to_sync.txt에도 추가 (이후 자동 동기화)
            add_result = add_channel_to_github(channel)
            if add_result == "already_exists":
                channel_msg = "📋 채널 목록에 이미 등록됨"
            else:
                channel_msg = "📋 채널 목록에 추가됨 (이후 자동 동기화)"
            persist_msg = f"💾 GitHub 저장 완료 (인덱스 재빌드 중)\n{channel_msg}"
        else:
            persist_msg = "⚠️ 이번 세션에만 유효"

        await status.edit_text(
            f"✅ @{channel} 동기화 완료!\n"
            f"💬 {msg_count}개 메시지 → {len(new_chunks)}개 청크 추가\n"
            f"📚 총 {context.bot_data['bm25'].n}개 청크\n"
            f"{persist_msg}\n\n"
            f"이제 이 채널 내용을 질문할 수 있습니다."
        )
    except Exception as e:
        await status.edit_text(f"❌ {type(e).__name__}: {str(e)[:150]}")

async def on_question(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    if not text:
        return

    # Notion URL 감지
    if re.match(r'https?://(www\.)?notion\.so/', text) or re.match(r'https?://[a-zA-Z0-9-]+\.notion\.site/', text):
        await handle_notion_url(update, context, text)
        return

    # Telegram 채널 URL 감지 (https://t.me/channelname 또는 @channelname)
    tme_match = re.match(r'https?://t\.me/([a-zA-Z0-9_]+)', text)
    at_match = re.match(r'^@([a-zA-Z0-9_]+)$', text)
    if tme_match:
        await handle_telegram_channel(update, context, tme_match.group(1))
        return
    if at_match:
        await handle_telegram_channel(update, context, at_match.group(1))
        return

    # 일반 URL도 Jina로 처리
    if re.match(r'https?://', text) and ' ' not in text:
        await handle_notion_url(update, context, text)
        return

    # 질문 처리
    bm25: BM25 = context.bot_data.get("bm25")
    if bm25 is None or bm25.n == 0:
        await update.message.reply_text(
            "아직 검색할 자료가 없습니다.\n파일이나 참고할 채널 링크를 먼저 보내 주세요."
        )
        return

    if not ANTHROPIC_API_KEY:
        await update.message.reply_text("답변 서비스 설정을 확인해야 합니다. 관리자에게 문의해 주세요.")
        return

    status = await update.message.reply_text("관련 자료를 확인하고 있습니다.")
    try:
        result = answer(text, bm25)
        await status.edit_text(result)
    except Exception as e:
        print(f"질문 답변 실패: {type(e).__name__}")
        await status.edit_text("답변을 완료하지 못했습니다. 잠시 후 다시 질문해 주세요.")

def fetch_index_from_github() -> list:
    # raw URL 사용 - Contents API의 1MB 용량 제한 없음
    url = f"https://raw.githubusercontent.com/{GITHUB_REPO}/main/data/index.json"
    try:
        resp = http_requests.get(url, timeout=60)
        if resp.status_code == 200:
            return resp.json()
    except Exception as e:
        print(f"GitHub 인덱스 로드 오류: {e}")
    return []

async def reload_index(context: ContextTypes.DEFAULT_TYPE):
    chunks = fetch_index_from_github()
    if not chunks:
        return
    old_n = context.bot_data.get("bm25", BM25([])).n
    new_bm25 = BM25(chunks)
    context.bot_data["bm25"] = new_bm25
    if new_bm25.n != old_n:
        print(f"인덱스 자동 갱신: {old_n} → {new_bm25.n}개 청크")

async def on_reload(update: Update, context: ContextTypes.DEFAULT_TYPE):
    status = await update.message.reply_text("🔄 GitHub에서 인덱스 다시 로드 중...")
    old_n = context.bot_data.get("bm25", BM25([])).n
    chunks = fetch_index_from_github()
    if not chunks:
        await status.edit_text("❌ GitHub에서 인덱스를 가져올 수 없습니다.")
        return
    context.bot_data["bm25"] = BM25(chunks)
    await status.edit_text(
        f"✅ 인덱스 갱신 완료!\n"
        f"📊 {old_n}개 → {len(chunks)}개 청크"
    )

async def on_status(update: Update, context: ContextTypes.DEFAULT_TYPE):
    bm25: BM25 = context.bot_data.get("bm25")
    n = bm25.n if bm25 else 0
    await update.message.reply_text(
        f"📚 인덱스: {n}개 청크\n"
        f"📁 파일 업로드: PDF, Excel, Word, TXT 지원\n"
        f"🔗 노션 URL: 링크를 그대로 보내면 자동 저장\n"
        f"📡 텔레그램 채널: @채널명 또는 t.me/링크 보내면 자동 동기화\n"
        f"🔄 /reload: GitHub에서 최신 인덱스 강제 갱신\n"
        f"💬 질문: 텍스트 그대로 입력"
    )

async def post_init(application):
    # GitHub에서 최신 인덱스 로드 (로컬 파일보다 우선)
    chunks = fetch_index_from_github()
    if not chunks and INDEX_PATH.exists():
        with open(INDEX_PATH, encoding="utf-8") as f:
            try:
                chunks = json.load(f)
            except Exception:
                chunks = []
    application.bot_data["bm25"] = BM25(chunks) if chunks else BM25([])
    print(f"인덱스 로드 완료: {len(chunks)}개 청크")
    # 30분마다 자동 갱신
    application.job_queue.run_repeating(reload_index, interval=1800, first=1800)

def main():
    app = Application.builder().token(TELEGRAM_TOKEN).post_init(post_init).build()
    app.add_handler(CommandHandler("status", on_status))
    app.add_handler(CommandHandler("reload", on_reload))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, on_question))
    print("QA 봇 실행 중...")
    app.run_polling()

if __name__ == "__main__":
    main()
